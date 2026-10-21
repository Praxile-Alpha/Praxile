"""Optional document parsers executed in an isolated, bounded-time child."""
from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path


MEDIA = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
}


def inspect_document(path: Path, kind: str, cell_refs=()) -> dict:
    if kind not in MEDIA.values():
        return {"status": "unsupported", "reason": "unknown_document_format"}
    if kind != "pdf":
        from defusedxml.ElementTree import fromstring
        from defusedxml.common import DefusedXmlException
        with zipfile.ZipFile(path) as package:
            members = package.infolist()
            if len(members) > 2000 or sum(item.file_size for item in members) > 50_000_000:
                return {"status": "failed", "reason": "expanded_package_limit"}
            if len({item.filename for item in members}) != len(members):
                return {"status": "failed", "reason": "duplicate_package_member"}
            if any(item.flag_bits & 1 for item in members):
                return {"status": "unsupported", "reason": "encrypted_package"}
            if any("vbaproject" in item.filename.lower() for item in members):
                return {"status": "unsupported", "reason": "macro_content"}
            for member in members:
                if member.filename.endswith((".xml", ".rels")):
                    payload = package.read(member)
                    try:
                        fromstring(payload, forbid_dtd=True, forbid_entities=True, forbid_external=True)
                    except DefusedXmlException:
                        return {"status": "failed", "reason": "xml_entity_declarations_forbidden"}
    if kind == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
            requested = {}
            for ref in cell_refs:
                column, row = coordinate_from_string(ref["cell"])
                requested[(ref["sheet"], row, column_index_from_string(column))] = {
                    **ref, "status": "missing_sheet" if ref["sheet"] not in book.sheetnames else "observed", "value": None}
            cells = 0
            for sheet in book:
                # Do not trust producer-supplied dimensions to hide requested cells.
                sheet.reset_dimensions()
                for row_number, row in enumerate(sheet.iter_rows(), 1):
                    cells += len(row)
                    if cells > 1_000_000:
                        return {"status": "failed", "reason": "cell_limit"}
                    for column_number, cell in enumerate(row, 1):
                        observation = requested.get((sheet.title, row_number, column_number))
                        if observation is None:
                            continue
                        if cell.data_type in {"f", "e"}:
                            observation.update(status="formula_or_error_unverified", value=None)
                        elif cell.value is None or type(cell.value) in (int, float, bool) or (
                                isinstance(cell.value, str) and len(cell.value) <= 2000):
                            observation["value"] = cell.value
                        else:
                            observation.update(status="unsupported_value", value=None)
            details = {"sheets": book.sheetnames, "cells_visited": cells, "formulas_recalculated": False}
            if cell_refs:
                details["cell_values"] = list(requested.values())
        finally:
            book.close()
    elif kind == "docx":
        from docx import Document
        document = Document(path)
        headings = [paragraph.text.strip() for paragraph in document.paragraphs
                    if paragraph.style is not None and paragraph.style.style_id in
                    {"Heading" + str(level) for level in range(1, 10)}]
        if len(headings) > 1000 or sum(map(len, headings)) > 100_000:
            return {"status": "failed", "reason": "heading_metadata_limit"}
        details = {"paragraphs": len(document.paragraphs), "tables": len(document.tables), "headings": headings}
    elif kind == "pptx":
        from pptx import Presentation
        document = Presentation(path)
        details = {"slides": len(document.slides)}
    else:
        from pypdf import PdfReader
        with path.open("rb") as source:
            document = PdfReader(source, strict=True)
            if document.is_encrypted:
                return {"status": "unsupported", "reason": "encrypted_pdf"}
            if len(document.pages) > 2000:
                return {"status": "failed", "reason": "page_limit"}
            details = {"pages": len(document.pages)}
    return {"status": "passed", "opened": True, "structure": details,
            "rendered": False, "semantic_quality_verified": False}


def main() -> None:
    # A resource ceiling supplements timeout on Linux. Other platforms should
    # run untrusted binary documents in the owning benchmark container.
    if sys.platform.startswith("linux"):
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (1_000_000_000, 1_000_000_000))
    try:
        result = inspect_document(Path(sys.argv[1]), sys.argv[2], json.loads(sys.argv[3]) if len(sys.argv) > 3 else ())
    except ImportError:
        result = {"status": "unsupported", "reason": "install_praxile_documents_extra"}
    except Exception as exc:
        result = {"status": "failed", "reason": "document_parse_failed", "error_type": type(exc).__name__}
    print(json.dumps(result))


if __name__ == "__main__":
    main()
